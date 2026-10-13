"""Excel из снимка уже рассчитанного отчёта. Без БД и внешних API."""
from datetime import date
from io import BytesIO
from math import isfinite


def build_manager_results_xlsx(data):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    def number(value):
        if isinstance(value, bool) or not isinstance(value, (float, int)) or not isfinite(value):
            raise ValueError('Некорректное число в отчёте')
        return value

    first, last = date.fromisoformat(data['from']), date.fromisoformat(data['to'])
    rows, totals = data['rows'], data['totals']
    if not isinstance(rows, list) or len(rows) > 5000 or first > last:
        raise ValueError('Некорректные данные отчёта')
    compare = bool(data.get('previous_period'))
    columns = [('manager', 'Менеджер', 'text'), ('system', 'Система', 'text'), ('calls', 'Звонки', 'count')]
    for key, label in [('directed', 'Направлено'), ('checked_in', 'Заселено'), ('shift_1', '1 смена'), ('shift_5', '5 смена')]:
        columns.extend([(key, label, 'count'), (key, label + ': % от звонков', 'rate')])
        if compare:
            columns.append((key, label + ': изменение к предыдущему периоду', 'change'))
    columns.append(('rating', 'Рейтинг', 'rating'))
    wb = Workbook()
    ws = wb.active
    ws.title = 'Результативность'
    width = len(columns)
    ws.append(['Результативность менеджеров'])
    ws.append([f'Период: {first:%d.%m.%Y} — {last:%d.%m.%Y}'])
    ws.append(['Все проценты считаются от звонков. Заселения и смены относятся к периоду направления.'])
    previous = data.get('previous_period')
    ws.append([f"Сравнение: {date.fromisoformat(previous['from']):%d.%m.%Y} — {date.fromisoformat(previous['to']):%d.%m.%Y}" if previous else 'Без сравнения'])
    for row in range(1, 5):
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=width)
    ws.append([label for _, label, _ in columns])
    ratings = [number(row['rating']) for row in rows]
    low, high = (min(ratings), max(ratings)) if ratings else (0, 0)
    for index, row in enumerate(rows + [dict(totals, manager='Итого', system='')]):
        total = index == len(rows)
        r = index + 6
        for c, (key, _, kind) in enumerate(columns, 1):
            color = 'F5E8E8' if total else ('F0F6FF' if index % 2 else 'FFFFFF')
            fmt = '#,##0'
            if kind == 'text':
                value = str(row.get(key, ''))[:500]
            elif kind == 'rate':
                value = number(row['rates'][key])
                fmt = '0.00%'
                if not total and value < number(totals['rates'][key]):
                    color = 'F4CCCC'
            elif kind == 'change':
                value = row.get('change', {}).get(key)
                fmt = '+0.0%;-0.0%;0.0%'
                if value is not None:
                    value = number(value)
                    if abs(value) >= 1e-8:
                        color = 'B7E1CD' if value > 0 else 'FCE5CD'
            elif kind == 'rating':
                value = None if total else number(row[key])
                fmt = '0.00'
                if not total:
                    fraction = (value - low) / (high - low) if high != low else 0
                    color = ''.join(f'{int(255 + (end - 255) * fraction + 0.5):02X}' for end in (230, 124, 115))
            else:
                value = number(row[key])
            cell = ws.cell(r, c, value)
            if kind == 'text':
                cell.data_type = 's'  # Имена никогда не превращаются в Excel-формулы.
            cell.number_format = fmt
            cell.fill = PatternFill('solid', fgColor=color)
            cell.font = Font(name='Calibri', size=11, bold=total)
            cell.alignment = Alignment(vertical='center', wrap_text=True)
        ws.row_dimensions[r].height = 34
    for cell in ws[5]:
        cell.fill = PatternFill('solid', fgColor='F5E8E8')
        cell.font = Font(bold=True)
        cell.alignment = Alignment(wrap_text=True, vertical='center')
    ws.row_dimensions[5].height = 65
    ws['A1'].font = Font(size=16, bold=True)
    for c, (_, _, kind) in enumerate(columns, 1):
        ws.column_dimensions[get_column_letter(c)].width = 42 if c == 1 else (22 if kind in ('rate', 'change') else 16)
    ws.freeze_panes = 'C6'
    ws.auto_filter.ref = f'A5:{get_column_letter(width)}{5 + len(rows)}'
    ws.sheet_view.showGridLines = False
    output = BytesIO()
    wb.save(output)
    output.seek(0)
    return output, f'manager_results_{first}_{last}.xlsx'
