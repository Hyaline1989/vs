"""Отчёты VIP, доступные администраторам с отдельным разрешением."""

import hmac
import logging
import os
from datetime import date
from io import BytesIO

from flask import Blueprint, jsonify, request, send_file
from flask_login import current_user

from report_refresh_store import finish_refresh
from shift_attendance_store import report_for_period, validate_snapshot, write_snapshot


service_reports_bp = Blueprint('service_reports', __name__, url_prefix='/api/service')
logger = logging.getLogger(__name__)


class InvalidPeriod(ValueError):
    pass


def _load_requested_shifts():
    """Одинаковая проверка периода для экрана и Excel."""
    requested_from = request.args.get('from', '')
    requested_to = request.args.get('to', '')
    try:
        start = date.fromisoformat(requested_from)
        end = date.fromisoformat(requested_to)
        if (start.isoformat() != requested_from or end.isoformat() != requested_to or
                not 2020 <= start.year <= 2099 or not 2020 <= end.year <= 2099):
            raise ValueError
    except ValueError:
        raise InvalidPeriod('Некорректный период') from None
    if end < start:
        raise InvalidPeriod('Дата «По» раньше даты «С»')
    if (end - start).days > 365:
        raise InvalidPeriod('Выберите период не более 366 дней')
    return report_for_period(start, end)


def _can_view_shifts():
    if not current_user.is_authenticated:
        return jsonify(error='Unauthorized'), 401
    if not current_user.can_view_vip():
        return jsonify(error='Forbidden'), 403
    return None


@service_reports_bp.get('/shifts')
def get_shifts():
    denied = _can_view_shifts()
    if denied is not None:
        return denied
    try:
        report = _load_requested_shifts()
    except InvalidPeriod as error:
        return jsonify(error=str(error)), 400
    except (OSError, ValueError) as error:
        logger.error('Не удалось прочитать сводку смен за %s — %s: %s',
                     request.args.get('from'), request.args.get('to'), error)
        return jsonify(error='Не удалось прочитать сохранённый отчёт'), 500
    return jsonify(report)


@service_reports_bp.get('/shifts.xlsx')
def export_shifts_xlsx():
    denied = _can_view_shifts()
    if denied is not None:
        return denied
    try:
        report = _load_requested_shifts()
    except InvalidPeriod as error:
        return jsonify(error=str(error)), 400
    except (OSError, ValueError) as error:
        logger.error('Не удалось прочитать сводку смен для Excel: %s', error)
        return jsonify(error='Не удалось прочитать сохранённый отчёт'), 500
    if report['status'] == 'not_loaded':
        return jsonify(error='Сводки за выбранный период ещё не загружены'), 404

    from openpyxl import Workbook
    from openpyxl.comments import Comment
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = 'Смены по табелям'
    sheet.freeze_panes = 'B2'
    sheet.append(['Объект', *(date.fromisoformat(day) for day in report['dates'])])
    sheet.column_dimensions['A'].width = 34
    header_fill = PatternFill('solid', fgColor='E9DCE0')
    stripe_fill = PatternFill('solid', fgColor='EDF5FF')
    missing_fill = PatternFill('solid', fgColor='FFF1CF')
    reason_labels = {
        'sheet_missing': 'В табеле нет листа за выбранный месяц',
        'sheet_ambiguous': 'В табеле несколько подходящих листов за месяц',
        'layout_unknown': 'Не удалось распознать ФИО или календарь',
        'calendar_incomplete': 'Этой даты нет в шапке календаря',
        'source_unavailable': 'Не удалось открыть табель',
        'month_not_loaded': 'Сводка за этот месяц ещё не загружена',
        'object_missing': 'Объект отсутствует в сводке за этот месяц',
    }
    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal='center')
    for index, day in enumerate(report['dates'], start=2):
        sheet.cell(1, index).number_format = 'dd.mm.yyyy'
        sheet.column_dimensions[get_column_letter(index)].width = 13
    for row_index, item in enumerate(report['objects'], start=2):
        name_cell = sheet.cell(row_index, 1, item['name'])
        name_cell.data_type = 's'  # Название из табеля не должно стать формулой Excel.
        for day_index, count in enumerate(item['counts']):
            cell = sheet.cell(row_index, day_index + 2, count if count is not None else '—')
            cell.alignment = Alignment(horizontal='center')
            if count is None:
                cell.fill = missing_fill
                cell.comment = Comment(reason_labels.get(item['reasons'][day_index], 'Нет данных'), 'Вектор')
            elif row_index % 2 == 0:
                cell.fill = stripe_fill
        if row_index % 2 == 0:
            name_cell.fill = stripe_fill
    total_row = len(report['objects']) + 2
    sheet.cell(total_row, 1, 'Итого')
    for cell in sheet[total_row]:
        cell.fill = header_fill
        cell.font = Font(bold=True)
    for day_index, total in enumerate(report['totals']):
        cell = sheet.cell(total_row, day_index + 2, total if total is not None else '—')
        cell.alignment = Alignment(horizontal='center')
        if report['incomplete'][day_index]:
            cell.fill = missing_fill
            cell.comment = Comment('Неполный итог: нет данных по одному или нескольким объектам', 'Вектор')
    sheet.sheet_view.showGridLines = False
    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    filename = f"smeny_po_tabelyam_{report['from']}_{report['to']}.xlsx"
    return send_file(output, as_attachment=True, download_name=filename,
                     mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')


@service_reports_bp.post('/shifts/import')
def import_shifts():
    expected = os.environ.get('PLANNED_DEPARTURES_IMPORT_KEY', '')
    if len(expected) < 32:
        return jsonify(error='Приём отчёта не настроен'), 503
    provided = request.headers.get('X-Vector-Import-Key', '')
    if not hmac.compare_digest(provided.encode(), expected.encode()):
        return jsonify(error='Unauthorized'), 401
    if request.content_length is not None and request.content_length > 256 * 1024:
        return jsonify(error='Слишком большой отчёт'), 413
    try:
        snapshot = validate_snapshot(request.get_json(silent=True))
        write_snapshot(snapshot)
    except ValueError as error:
        return jsonify(error=str(error)), 400
    except OSError as error:
        logger.error('Не удалось сохранить сводку смен: %s', error)
        return jsonify(error='Не удалось сохранить отчёт'), 500
    refresh_id = request.headers.get('X-Vector-Refresh-Id', '')
    if refresh_id:
        try:
            finish_refresh('shifts', snapshot['month'], refresh_id, 'completed')
        except OSError as error:
            logger.error('Не удалось отметить обновление сводки смен: %s', error)
    return jsonify(status='saved', month=snapshot['month'], objects=len(snapshot['objects']))
