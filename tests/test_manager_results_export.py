import unittest
from openpyxl import load_workbook
from manager_results_export import build_manager_results_xlsx


class ExcelExportTests(unittest.TestCase):
    def report(self):
        row = dict(manager='=Тест', system='Комплект', calls=10, directed=2,
                   checked_in=1, shift_1=1, shift_5=0, rating=1.5,
                   rates=dict(directed=.2, checked_in=.1, shift_1=.1, shift_5=0),
                   change=dict(directed=.5, checked_in=None, shift_1=-.2, shift_5=0))
        return {'from': '2026-10-01', 'to': '2026-10-07', 'rows': [row],
                'totals': dict(row), 'previous_period': None}

    def test_numeric_values_styles_and_text_safety(self):
        output, name = build_manager_results_xlsx(self.report())
        sheet = load_workbook(output).active
        self.assertEqual(sheet['A6'].value, '=Тест')
        self.assertEqual(sheet['A6'].data_type, 's')
        self.assertEqual(sheet['C6'].value, 10)
        self.assertEqual(sheet['E6'].value, .2)
        self.assertEqual(sheet['E6'].number_format, '0.00%')
        self.assertEqual(sheet['A7'].value, 'Итого')
        self.assertEqual(sheet.freeze_panes, 'C6')
        self.assertTrue(name.endswith('.xlsx'))

    def test_comparison_and_empty_report(self):
        data = self.report()
        data['previous_period'] = {'from': '2026-09-24', 'to': '2026-09-30'}
        output, _ = build_manager_results_xlsx(data)
        sheet = load_workbook(output).active
        self.assertEqual(sheet['F6'].value, .5)
        self.assertEqual(sheet['F6'].fill.fgColor.rgb, '00B7E1CD')
        data['rows'] = []
        output, _ = build_manager_results_xlsx(data)
        self.assertEqual(load_workbook(output).active['A6'].value, 'Итого')

    def test_bad_number_rejected(self):
        data = self.report()
        data['rows'][0]['calls'] = '=1+2'
        with self.assertRaises(ValueError):
            build_manager_results_xlsx(data)
